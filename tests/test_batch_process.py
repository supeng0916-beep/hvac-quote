"""批量检查图纸 / 导入模板价格的后台任务与进度。"""

import io
import time

import ezdxf
from openpyxl import Workbook
from fastapi.testclient import TestClient

from hvac_quote.app import create_app


def _client(tmp_path) -> tuple:
    app = create_app(runs_dir=tmp_path / "runs", data_dir=tmp_path / "data")
    client = TestClient(app)
    client.post(
        "/setup",
        data={
            "display_name": "管理员",
            "username": "admin",
            "password": "password123",
            "password_confirm": "password123",
        },
    )
    return app, client


def _project(client: TestClient) -> str:
    response = client.post(
        "/projects",
        data={"name": "批量处理项目", "customer": "测试客户"},
        follow_redirects=False,
    )
    assert response.status_code == 303
    return response.headers["location"].rsplit("/", 1)[-1]


def _upload(client: TestClient, project_id: str, name: str, content: bytes, kind: str):
    response = client.post(
        f"/projects/{project_id}/files",
        data={"kind": kind},
        files={"upload": (name, content, "application/octet-stream")},
        follow_redirects=False,
    )
    assert response.status_code == 303


def _template_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "空调"
    headers = ["名称", "规格及型号", "单位", "工程量", "材料单价/元", "安装费单价/元"]
    values = ["室内机", "DEMO-UNIT-01", "台", 1, 1200, 100]
    for column, (header, value) in enumerate(zip(headers, values), start=2):
        sheet.cell(1, column, header)
        sheet.cell(2, column, value)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _dxf_bytes() -> bytes:
    doc = ezdxf.new("R2010")
    buffer = io.StringIO()
    doc.write(buffer, fmt="asc")
    return buffer.getvalue().encode("utf-8")


def _wait_batch_done(client: TestClient, project_id: str, timeout: float = 20.0):
    deadline = time.monotonic() + timeout
    page = None
    while time.monotonic() < deadline:
        page = client.get(f"/projects/{project_id}/files")
        if page.status_code == 200 and "批量处理完成" in page.text:
            return page
        time.sleep(0.2)
    return page


def test_batch_processes_pending_files(tmp_path):
    app, client = _client(tmp_path)
    project_id = _project(client)
    _upload(client, project_id, "模板.xlsx", _template_bytes(), "template")
    _upload(client, project_id, "一层.dxf", _dxf_bytes(), "drawing")

    started = client.post(
        f"/projects/{project_id}/files/batch-process",
        follow_redirects=False,
    )
    assert started.status_code == 303
    assert "batch-started" in started.headers["location"]

    page = _wait_batch_done(client, project_id)
    assert page is not None and "批量处理完成" in page.text
    assert "成功 2" in page.text

    assert len(app.state.new_quotes.list_inspections(project_id)) == 1
    assert len(app.state.new_quotes.list_price_books(project_id)) == 1

    actions = {
        str(row["action"])
        for row in app.state.db.query("SELECT action FROM audit_logs")
    }
    assert {"batch_run", "template_imported", "drawing_inspected"} <= actions


def test_batch_rejects_second_run_while_running(tmp_path):
    app, client = _client(tmp_path)
    project_id = _project(client)
    app.state.batch_jobs[project_id] = {
        "status": "running",
        "total": 1,
        "done": 0,
        "current": "一层.dxf",
        "startedAt": "2026-10-21T00:00:00",
        "finishedAt": "",
        "results": [],
    }

    response = client.post(
        f"/projects/{project_id}/files/batch-process",
        follow_redirects=False,
    )
    assert response.status_code == 303
    assert "batch-running" in response.headers["location"]
    page = client.get(response.headers["location"])
    assert "已有批量任务" in page.text
    assert app.state.batch_jobs[project_id]["status"] == "running"


def test_batch_records_file_errors_and_continues(tmp_path, monkeypatch):
    app, client = _client(tmp_path)
    project_id = _project(client)
    _upload(client, project_id, "坏图.dxf", _dxf_bytes(), "drawing")
    import hvac_quote.app as app_module

    original_inspect = app_module.inspect_drawing
    calls = {"count": 0}

    def broken_inspect(path, file_id, workdir):
        calls["count"] += 1
        if calls["count"] == 1:
            raise ValueError("读取失败（测试注入）")
        return original_inspect(path, file_id, workdir)

    monkeypatch.setattr(app_module, "inspect_drawing", broken_inspect)
    _upload(client, project_id, "模板.xlsx", _template_bytes(), "template")

    started = client.post(
        f"/projects/{project_id}/files/batch-process",
        follow_redirects=False,
    )
    assert started.status_code == 303

    page = _wait_batch_done(client, project_id)
    assert page is not None and "批量处理完成" in page.text
    assert "成功 1" in page.text
    assert "失败 1" in page.text
    assert "坏图.dxf" in page.text

    assert len(app.state.new_quotes.list_inspections(project_id)) == 0
    assert len(app.state.new_quotes.list_price_books(project_id)) == 1


def test_batch_empty_when_nothing_pending(tmp_path):
    app, client = _client(tmp_path)
    project_id = _project(client)
    _upload(client, project_id, "参考.pdf", b"%PDF-1.4", "reference")

    response = client.post(
        f"/projects/{project_id}/files/batch-process",
        follow_redirects=False,
    )
    assert response.status_code == 303
    assert "batch-empty" in response.headers["location"]
    page = client.get(response.headers["location"])
    assert "没有需要处理" in page.text
