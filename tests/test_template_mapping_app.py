from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook

from hvac_quote.app import create_app


def _workbook(path: Path, *, sheet: str, header_row: int, columns: dict, row: int):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    semantic = {
        "名称": "name",
        "规格及型号": "spec",
        "单位": "unit",
        "工程量": "quantity",
        "材料单价/元": "material",
        "安装费单价/元": "installation",
    }
    values = {
        "name": "风冷模块主机",
        "spec": "DEMO-CHILLER-01",
        "unit": "台",
        "quantity": 5,
        "material": 12300,
        "installation": 300,
    }
    for header, column in columns.items():
        worksheet.cell(header_row, column, header)
        worksheet.cell(row, column, values[semantic[header]])
    workbook.save(path)


def _client(tmp_path: Path) -> TestClient:
    app = create_app(runs_dir=tmp_path / "runs", data_dir=tmp_path / "data")
    client = TestClient(app)
    client.post(
        "/setup",
        data={
            "display_name": "映射管理员",
            "username": "admin",
            "password": "password123",
            "password_confirm": "password123",
        },
    )
    return client


def test_manual_mapping_is_saved_and_applied_to_next_workbook(tmp_path: Path):
    client = _client(tmp_path)
    project_id = client.post(
        "/projects",
        data={"name": "客户模板映射测试", "customer": "测试客户"},
        follow_redirects=False,
    ).headers["location"].rsplit("/", 1)[-1]

    first = tmp_path / "first.xlsx"
    _workbook(
        first,
        sheet="客户报价A",
        header_row=3,
        columns={
            "名称": 2,
            "规格及型号": 3,
            "单位": 4,
            "工程量": 11,
            "材料单价/元": 12,
            "安装费单价/元": 13,
        },
        row=4,
    )
    client.post(
        f"/projects/{project_id}/files",
        data={"kind": "template"},
        files={
            "upload": (
                "first.xlsx",
                first.read_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        follow_redirects=False,
    )
    first_file = next(
        item
        for item in client.app.state.projects.get("admin", project_id).files
        if item.original_name == "first.xlsx"
    )
    mapped = client.post(
        f"/projects/{project_id}/files/{first_file.id}/import-template",
        data={
            "sheet": "客户报价A",
            "header_row": "3",
            "name_col": "B",
            "spec_col": "C",
            "unit_col": "D",
            "quantity_col": "K",
            "material_col": "L",
            "install_col": "M",
        },
        follow_redirects=False,
    )
    assert mapped.status_code == 303
    profiles = client.app.state.mapping_profiles.list_for(
        "admin", "name:测试客户"
    )
    assert len(profiles) == 1
    assert profiles[0]["sheetName"] == "客户报价A"

    second = tmp_path / "second.xlsx"
    _workbook(
        second,
        sheet="2027报价",
        header_row=2,
        columns={
            "名称": 3,
            "规格及型号": 4,
            "单位": 5,
            "工程量": 8,
            "材料单价/元": 9,
            "安装费单价/元": 10,
        },
        row=3,
    )
    client.post(
        f"/projects/{project_id}/files",
        data={"kind": "template"},
        files={
            "upload": (
                "second.xlsx",
                second.read_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        follow_redirects=False,
    )
    second_file = next(
        item
        for item in client.app.state.projects.get("admin", project_id).files
        if item.original_name == "second.xlsx"
    )
    imported = client.post(
        f"/projects/{project_id}/files/{second_file.id}/import-template",
        follow_redirects=False,
    )
    assert imported.status_code == 303
    books = client.app.state.new_quotes.list_price_books(project_id)
    second_book = next(
        item for item in books if item["fileId"] == second_file.id
    )
    assert second_book["book"]["detectedBy"] == "profile"
    assert second_book["book"]["counts"]["priced"] == 1
    page = client.get(imported.headers["location"])
    assert "该客户上次保存的列映射" in page.text

    list_page = client.get(f"/projects/{project_id}/template-mappings")
    assert list_page.status_code == 200
    assert "客户模板映射" in list_page.text
    assert "first.xlsx" in list_page.text
    profile_id = client.app.state.mapping_profiles.list_for(
        "admin", "name:测试客户"
    )[0]["id"]
    deleted = client.post(
        f"/projects/{project_id}/template-mappings/{profile_id}/delete",
        follow_redirects=False,
    )
    assert deleted.status_code == 303
    assert client.app.state.mapping_profiles.list_for(
        "admin", "name:测试客户"
    ) == []
    after = client.get(f"/projects/{project_id}/template-mappings?notice=deleted")
    assert "已删除" in after.text
