from pathlib import Path

from openpyxl import Workbook

from hvac_quote.template_book import (
    import_price_book_auto,
    import_price_book_mapped,
    inspect_workbook_tables,
)


def test_auto_import_finds_nonstandard_sheet_and_columns(tmp_path: Path):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "汇总"
    summary["A1"] = "项目汇总"
    summary["A2"] = "这里不是报价明细"

    quote = workbook.create_sheet("示例报价表")
    quote["B3"] = "名称"
    quote["C3"] = "规格及型号"
    quote["D3"] = "单位"
    quote["K3"] = "工程量"
    quote["L3"] = "材料单价/元"
    quote["M3"] = "安装费单价/元"
    quote.append(
        [None, "风冷模块主机", "DEMO-CHILLER-01", "台", None, None, None, None, None, None, 5, 12300, 300]
    )
    quote.append(
        [None, "镀锌钢管", "DN25*2.75mm", "米", None, None, None, None, None, None, 35, 12.50, 50]
    )
    path = tmp_path / "synthetic-quote.xlsx"
    workbook.save(path)

    book = import_price_book_auto(
        path, file_id="template", source_sha256="x"
    )

    assert book["sheetName"] == "示例报价表"
    assert book["headerRow"] == 3
    assert book["detectedBy"] == "auto"
    assert book["columns"]["name"] == "B"
    assert book["columns"]["quantity"] == "K"
    assert book["columns"]["material"] == "L"
    assert book["columns"]["installation"] == "M"
    assert book["counts"]["priced"] == 2
    assert book["rows"][0]["material"] == "12300"
    assert book["rows"][0]["installation"] == "300"
    assert book["rows"][0]["category"] == "equipment"
    assert book["rows"][1]["category"] == "pipe"


def test_manual_mapping_imports_selected_sheet_and_columns(tmp_path: Path):
    workbook = Workbook()
    workbook.active.title = "说明"
    workbook.active["A1"] = "说明页"
    quote = workbook.create_sheet("客户报价")
    quote["B3"] = "名称"
    quote["C3"] = "规格及型号"
    quote["D3"] = "单位"
    quote["K3"] = "工程量"
    quote["L3"] = "材料单价/元"
    quote["M3"] = "安装费单价/元"
    quote.append([None, "风冷模块主机", "DEMO-CHILLER-01", "台", None, None, None, None, None, None, 5, 12300, 300])
    path = tmp_path / "manual.xlsx"
    workbook.save(path)

    tables = inspect_workbook_tables(path)
    assert [table["name"] for table in tables] == ["说明", "客户报价"]
    assert any(table["headerRow"] == 3 for table in tables)

    book = import_price_book_mapped(
        path,
        file_id="template",
        source_sha256="x",
        sheet_name="客户报价",
        header_row=3,
        columns={
            "name": "B",
            "spec": "C",
            "unit": "D",
            "quantity": "K",
            "material": "L",
            "installation": "M",
        },
    )
    assert book["detectedBy"] == "manual"
    assert book["sheetName"] == "客户报价"
    assert book["columns"]["material"] == "L"
    assert book["counts"]["priced"] == 1
    assert book["rows"][0]["installation"] == "300"
