from pathlib import Path

from openpyxl import Workbook

from hvac_quote.db import Database
from hvac_quote.mapping_profiles import MappingProfileStore
from hvac_quote.template_book import (
    find_matching_mapping,
    import_price_book_mapped,
)


def _workbook(path: Path, *, sheet: str, header_row: int, columns: dict, row: int):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    for key, column in columns.items():
        worksheet.cell(header_row, column, key)
    values = {
        "name": "风冷模块主机",
        "spec": "DEMO-CHILLER-01",
        "unit": "台",
        "quantity": 5,
        "material": 12300,
        "installation": 300,
    }
    semantic = {
        "名称": "name",
        "规格及型号": "spec",
        "单位": "unit",
        "工程量": "quantity",
        "材料单价/元": "material",
        "安装费单价/元": "installation",
    }
    for key, column in columns.items():
        worksheet.cell(row, column, values[semantic[key]])
    workbook.save(path)


def test_saved_mapping_applies_to_same_header_layout(tmp_path: Path):
    db = Database(tmp_path / "test.db")
    store = MappingProfileStore(db)
    first = tmp_path / "first.xlsx"
    _workbook(
        first,
        sheet="客户报价A",
        header_row=3,
        columns={
            "名称": 2,
            "规格及型号": 3,
            "单位": 4,
            "工程量": 11,
            "材料单价/元": 12,
            "安装费单价/元": 13,
        },
        row=4,
    )
    book1 = import_price_book_mapped(
        first,
        file_id="one",
        source_sha256="x",
        sheet_name="客户报价A",
        header_row=3,
        columns={
            "name": "B",
            "spec": "C",
            "unit": "D",
            "quantity": "K",
            "material": "L",
            "installation": "M",
        },
    )
    profile = store.upsert(
        owner="admin",
        key="name:测试客户",
        signature=book1["mappingSignature"],
        sheet_name=book1["sheetName"],
        header_row=book1["headerRow"],
        columns=book1["columns"],
        headers=book1["headerLabels"],
        file_name="first.xlsx",
    )
    assert profile["id"].startswith("map-")

    second = tmp_path / "second.xlsx"
    _workbook(
        second,
        sheet="2027报价",
        header_row=2,
        columns={
            "名称": 3,
            "规格及型号": 4,
            "单位": 5,
            "工程量": 8,
            "材料单价/元": 9,
            "安装费单价/元": 10,
        },
        row=3,
    )
    matched = find_matching_mapping(
        second, store.list_for("admin", "name:测试客户")
    )
    assert matched is not None
    assert matched["sheetName"] == "2027报价"
    assert matched["columns"] == {
        "name": "C",
        "spec": "D",
        "unit": "E",
        "quantity": "H",
        "material": "I",
        "installation": "J",
    }

    book2 = import_price_book_mapped(
        second,
        file_id="two",
        source_sha256="y",
        sheet_name=matched["sheetName"],
        header_row=matched["headerRow"],
        columns=matched["columns"],
    )
    assert book2["counts"]["priced"] == 1
    assert book2["rows"][0]["material"] == "12300"
    assert book2["rows"][0]["installation"] == "300"


def test_profile_can_be_deleted(tmp_path: Path):
    db = Database(tmp_path / "delete.db")
    store = MappingProfileStore(db)
    profile = store.upsert(
        owner="admin",
        key="name:测试客户",
        signature="sig",
        sheet_name="报价",
        header_row=3,
        columns={"name": "B", "spec": "C", "unit": "D", "material": "G"},
        headers={"name": "名称", "spec": "规格", "unit": "单位", "material": "材料单价"},
        file_name="x.xlsx",
    )
    assert store.get("admin", "name:测试客户", profile["id"]) is not None
    assert store.delete("admin", "name:测试客户", profile["id"]) is True
    assert store.get("admin", "name:测试客户", profile["id"]) is None
    assert store.delete("admin", "name:测试客户", profile["id"]) is False
