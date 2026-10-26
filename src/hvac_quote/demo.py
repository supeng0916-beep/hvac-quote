"""Reproducible, wholly synthetic demo. Requires the development dependencies."""
from __future__ import annotations

import hashlib
import json
import secrets
from pathlib import Path


def generate_samples(directory: Path, *, alternate_layout: bool = False) -> tuple[Path, Path]:
    """Generate new geometry and fictitious prices; never read customer inputs."""
    import ezdxf
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    directory.mkdir(parents=True, exist_ok=True)
    drawing, template = directory / 'demo-office.dxf', directory / 'demo-prices.xlsx'
    if drawing.exists() or template.exists():
        raise FileExistsError('示例文件已存在；请选择新的目录。')
    doc = ezdxf.new('R2018')
    doc.header['$INSUNITS'] = 4
    for name in ['HVAC设备', '冷媒铜管', '冷凝水 DN25', '示例说明']:
        doc.layers.new(name)
    block = doc.blocks.new('KFR-26GW-DEMO')
    block.add_lwpolyline([(0,0),(600,0),(600,400),(0,400)], close=True)
    model = doc.modelspace()
    for point in [(0,0),(8000,0)]:
        model.add_blockref('KFR-26GW-DEMO', point, dxfattribs={'layer':'HVAC设备'})
    model.add_line((0,1500),(8000,1500),dxfattribs={'layer':'冷媒铜管','color':1})
    model.add_text('9.53x15.88',dxfattribs={'insert':(3000,1550),'height':100})
    model.add_line((0,4000),(6000,4000),dxfattribs={'layer':'冷凝水 DN25','color':5})
    model.add_text('DN25',dxfattribs={'insert':(2500,4050),'height':100})
    model.add_text('SYNTHETIC DEMO - NOT FOR CONSTRUCTION',dxfattribs={'layer':'示例说明','insert':(0,5500),'height':200})
    doc.saveas(drawing)
    wb = Workbook(); ws = wb.active; ws.title = '演示报价'
    wb.properties.creator = 'HVAC Quote Demo'; wb.properties.lastModifiedBy = 'HVAC Quote Demo'
    ws['A1'] = '虚构办公空间 · 软件演示价格表'
    ws['A2'] = '全部图纸、数量和价格为人工构造，仅用于软件演示。'
    headers=['序号','名称','规格型号','单位','备注','数量','材料单价','安装单价','说明','材料合计','安装合计','合计']
    ws.append([])
    for col,name in enumerate(headers,1):
        c=ws.cell(5,col,name);c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='18334A')
    rows={6:('示例分体空调','KFR-26GW-DEMO','台',1200,100),101:('铜管','9.53','米',20,4),102:('铜管','15.88','米',30,6),103:('冷凝水排水管','DN25','米',5,2)}
    for index,(row,(name,spec,unit,material,installation)) in enumerate(rows.items(),1):
        for col,value in {1:index,2:name,3:spec,4:unit,6:0,7:material,8:installation,10:f'=F{row}*G{row}',11:f'=F{row}*H{row}',12:f'=J{row}+K{row}'}.items(): ws.cell(row,col,value)
    for cell,formula in {'J98':'=SUM(J6:J97)','J99':'=J98+SUM(K6:K97)','J162':'=SUM(J100:J161)','K162':'=SUM(K100:K161)','L162':'=J162+K162','L163':'=L162','L164':'=J98+SUM(K6:K97)+L162','L165':'=L164'}.items():ws[cell]=formula
    for col in 'BC':ws.column_dimensions[col].width=28
    for col in 'FGHJKL':ws.column_dimensions[col].width=15
    ws.freeze_panes='F6'
    if alternate_layout:
        wb.remove(ws)
        cover=wb.create_sheet('说明');cover['A1']='Synthetic demo only'
        ws=wb.create_sheet('自定义价格表')
        for col,label in {2:'名称',3:'规格型号',4:'单位',11:'工程量',12:'材料单价',13:'安装单价'}.items():ws.cell(3,col,label)
        for row,values in enumerate(rows.values(),4):
            name,spec,unit,material,installation=values
            for col,value in {2:name,3:spec,4:unit,11:0,12:material,13:installation}.items():ws.cell(row,col,value)
    wb.save(template)
    return drawing,template


def create_demo(directory: Path, *, alternate_layout: bool = False) -> dict:
    """Build fresh isolated demo state via real application request handlers."""
    from fastapi.testclient import TestClient
    from .app import create_app
    directory=Path(directory).resolve()
    if directory.exists() and any(directory.iterdir()):
        raise FileExistsError('演示目录必须为空；不会覆盖已有账号或项目。')
    drawing,template=generate_samples(directory/'inputs', alternate_layout=alternate_layout)
    original={p:hashlib.sha256(p.read_bytes()).hexdigest() for p in (drawing,template)}
    password=secrets.token_urlsafe(18)
    app=create_app(directory/'runs',data_dir=directory/'data')
    with TestClient(app) as client:
        def post(path,**kwargs):
            response=client.post(path,follow_redirects=False,**kwargs)
            if response.status_code != 303:
                raise RuntimeError(f'{path}: HTTP {response.status_code}\n{response.text[:4000]}')
            return response.headers['location']
        post('/setup',data={'username':'demo','display_name':'演示用户','password':password,'password_confirm':password})
        project_path=post('/projects',data={'name':'示例办公空间 · 暖通安装','customer':'虚构演示客户','note':'人工构造数据，仅用于软件演示，不代表真实工程或市场价格。'})
        project_id=project_path.rsplit('/',1)[-1]
        for kind,path in [('drawing',drawing),('template',template)]:
            post(project_path+'/files',data={'kind':kind},files={'upload':(path.name,path.read_bytes())})
        project=app.state.projects.get('demo',project_id)
        for item in project.files:
            suffix='inspect' if item.kind=='drawing' else 'import-template'
            post(f'{project_path}/files/{item.id}/{suffix}',data={'views':'model'} if item.kind=='drawing' else {})
        inspection=app.state.new_quotes.list_inspections(project_id)[0]
        book=app.state.new_quotes.list_price_books(project_id)[0]
        draft_path=post(project_path+'/custom/drafts',data={'inspection_id':inspection['id'],'price_book_id':book['id'],'title':'演示报价 · 已核对范围'})
        draft_id=draft_path.rsplit('/',1)[-1]
        draft=app.state.new_quotes.get_draft(project_id,draft_id);data=draft['data']
        form={'expected_revision':str(draft['revision']),'scale':'0.001','action':'generate','taxRate':'0','globalCoefficient':'1'}
        for i,g in enumerate(data['equipment']):
            form.update({f'equipment_include_{i}':'on',f'equipment_quantity_{i}':str(g['quantity']),f'equipment_price_{i}':str(next(row['rowNumber'] for row in book['book']['rows'] if row['category']=='equipment')),f'equipment_set_confirm_{i}':'on'})
        for i,g in enumerate(data['routes']):
            row=next(row['rowNumber'] for row in book['book']['rows'] if str(row['spec']) == ('DN25' if str(g['spec']) in {'25','DN25'} else str(g['spec'])))
            form.update({f'route_include_{i}':'on',f'route_length_{i}':str(g['lengthM']),f'route_price_{i}':str(row),f'route_review_confirm_{i}':'on'})
        for i in range(len(data.get('issues') or [])): form[f'ack_issue_{i}']='on'
        post(draft_path,data=form)
        draft=app.state.new_quotes.get_draft(project_id,draft_id)
        project=app.state.projects.get('demo',project_id)
        if draft['status'] != 'ready' or not project.quotes:
            raise RuntimeError('演示报价未生成：'+json.dumps(draft.get('exceptions'),ensure_ascii=False))
        quote=project.quotes[0]
        result_path=f'{project_path}/quotes/{quote.run_id}'
        assert client.get(result_path).status_code==200
        result={'username':'demo','password':password,'project':project_path,'inspection':f"{project_path}/inspections/{inspection['id']}",'price_book':f"{project_path}/price-books/{book['id']}",'draft':draft_path,'result':result_path,'total':quote.total,'excel':f'runs/{quote.run_id}/quote.xlsx','pdf':f'runs/{quote.run_id}/quote.pdf','equipment_quantity':sum(int(g['quantity']) for g in data['equipment']),'lengths':[float(g['lengthM']) for g in data['routes']],'inputs_unchanged':all(hashlib.sha256(p.read_bytes()).hexdigest()==v for p,v in original.items())}
        (directory/'preview.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
        return result


def main():
    import argparse
    parser=argparse.ArgumentParser(description='创建独立的虚构演示项目；需要 uv sync 的开发依赖。')
    parser.add_argument('--out',type=Path,default=Path('data/demo'))
    args=parser.parse_args()
    result=create_demo(args.out)
    print('演示项目已创建。所有资料和价格均为虚构。')
    print(f"账号：{result['username']}\n随机密码：{result['password']}")
    print(f'启动：uv run hvac-quote serve --data "{args.out}/data" --runs "{args.out}/runs"')

if __name__=='__main__':main()
